"""Supplementary tables S1-S10 as one Excel workbook.

S1  Anchor studies and anchor accessions
S2  Metadata and quality control for every record (accession versions)
S3  Haplotype, cluster, nearest neighbours and anchor verdict for every
    record that passed quality control
S4  Sensitivity: K2P thresholds, anchor evidence and taxonomic treatment
S5  Identification test: every query, outcomes by species, raw -> cleaned
S6  Records disagreeing with the anchored species of their cluster, and
    COI-ITS2 comparison for specimens sequenced for both genes
S7  Tips of the four trees, in figure order, with the records behind each tip
S8  Oligonychus: anchor studies and anchor accessions
S9  Oligonychus: records, clusters, identification test and sensitivity
S10 Change by submission period and country, both genera

Usage: python scripts/supplementary_tables.py [--out results/supplementary]
"""
import argparse
import csv
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


def read(path):
    if not os.path.exists(path):
        return []
    with open(path, newline="") as f:
        return list(csv.DictReader(f))


def clean(text):
    """Drop notes on where a detail was found (they belong in the data file, not the table)."""
    text = re.sub(r"\s*\((?:per|claimed by)[^)]*\)", "", text or "")
    text = re.sub(r"\s*claimed by [A-Za-z ,]+ only", "", text).strip()
    text = re.sub(r"^uncertain", "not stated", text)
    return text[:1].upper() + text[1:]


def add_sheet(wb, title, caption, blocks):
    ws = wb.create_sheet(title)
    ws.append([caption])
    ws["A1"].font = Font(bold=True)
    for heading, header, rows in blocks:
        ws.append([])
        if heading:
            ws.append([heading])
            ws.cell(ws.max_row, 1).font = Font(italic=True)
        ws.append(header)
        for c in range(1, len(header) + 1):
            ws.cell(ws.max_row, c).font = Font(bold=True)
        for r in rows:
            ws.append([r.get(h, "") if isinstance(r, dict) else r[i] for i, h in enumerate(header)])
    for col in range(1, ws.max_column + 1):
        width = max((len(str(ws.cell(r, col).value or "")) for r in range(2, ws.max_row + 1)), default=8)
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 8), 60)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top")
    ws.freeze_panes = None
    return ws


def anchor_tables(studies, anchors, status, merged):
    """Source and accession rows for an anchor table (S1, S8)."""
    # S1: details checked against the publication are given as found; anything not
    # checked against the full text is marked "unverified".
    def detail(s, key):
        if s["evidence_level"] not in "AB":
            return "Not verified (lead only)"
        value = clean(s.get(key, ""))
        if s.get("verification_status", "").startswith("verified"):
            return value or "Not stated"
        return f"{value} (unverified)" if value else "Unverified"

    s1_studies = []
    for s in studies:
        verified = s.get("verification_status", "").startswith("verified")
        s1_studies.append({
            "Source": s["citation"],
            "Source type": "GenBank direct submission (unpublished)" if "direct submission" in s["citation"].lower()
            else "Publication",
            "DOI": s["doi"] if s["doi"].startswith("10.") and "check" not in s["doi"] else "",
            "Evidence level": s["evidence_level"] + ("" if verified or s["evidence_level"] not in "AB"
                                                     else " (provisional)"),
            "Used for decisions": "yes" if s["evidence_level"] in "AB" else "no (lead only)",
            "Basis of identification": detail(s, "morphology_verified"),
            "Males examined": detail(s, "males_examined"), "Identifier": detail(s, "identifier_named"),
            "Voucher repository": detail(s, "voucher_linked"), "Crossing tests": detail(s, "crossing_tests"),
            "Material sequenced": (clean(s.get("material_sequenced", "")) or "Not stated") if verified else
            ("Unverified" if s["evidence_level"] in "AB" else "Not verified (lead only)"),
            "Supporting passage (section, page)": s.get("supporting_passage", "") if verified else "Unverified",
            "Checked against full text": "yes" if verified else "no",
            "Anchor accessions": sum(a["study_id"] == s["study_id"] for a in anchors)})
    by_id = {s["study_id"]: s for s in studies}
    s1_acc = []
    for a in anchors:
        st = status.get(a["accession"], {})
        rid = st.get("record_id", "")
        s1_acc.append({"Accession": rid or a["accession"], "Study": by_id[a["study_id"]]["citation"].split(" (")[0]
                       + " (" + by_id[a["study_id"]]["citation"].split(" (")[1].split(")")[0] + ")",
                       "Evidence level": st.get("evidence_level") or by_id[a["study_id"]]["evidence_level"],
                       "Name given in the study": a.get("species", ""),
                       "Name in record": merged.get(rid, {}).get("species_label", ""),
                       "Accepted name": st.get("species_name", ""),
                       "Passed QC": st.get("passed_qc", ""), "Cluster": st.get("cluster", ""),
                       "Anchored species of cluster": st.get("cluster_anchored_species", ""),
                       "Consistency": st.get("anchor_consistency", "")})
    h1 = list(s1_studies[0].keys())
    h1b = list(s1_acc[0].keys())
    return h1, s1_studies, h1b, s1_acc


def main():
    p = argparse.ArgumentParser(description="Supplementary tables")
    p.add_argument("--out", default=os.path.join("results", "supplementary"))
    args = p.parse_args()
    os.makedirs(args.out, exist_ok=True)

    studies = read("data/anchors/anchor_studies.csv")
    anchors = read("data/anchors/anchor_records.csv")
    status = {r["accession"]: r for r in read("results/anchors/anchor_status.csv")}
    merged = {r["record_id"]: r for r in read("results/genus/merged_records.csv")}
    qc = read("results/genus/qc_records.csv")
    checks = {r["record_id"]: r for r in read("results/species/record_checks.csv")}
    verdicts = {r["record_id"]: r for r in read("results/anchors/record_anchor_check.csv")}
    nn = {r["record_id"]: r for r in read("results/robustness/coi_nearest_neighbour.csv")}

    wb = Workbook()
    wb.remove(wb.active)

    studies = read("data/anchors/anchor_studies.csv")
    anchors = read("data/anchors/anchor_records.csv")
    h1, s1_studies, h1b, s1_acc = anchor_tables(studies, anchors, status, merged)
    add_sheet(wb, "S1 Anchors", "Table S1. Candidate anchor sources (eleven publications and one unpublished "
              "GenBank submission), evidence levels and COI accessions. Details not checked against the full "
              "text of the publication are marked unverified.",
              [("Sources", h1, s1_studies), ("Anchor accessions", h1b, s1_acc)])

    # S2
    s2 = []
    for r in qc:
        m = merged.get(r["record_id"], {})
        s2.append({"Record": r["record_id"], "GenBank accession": m.get("genbank_accession", ""),
                   "BOLD process ID": m.get("bold_processid", ""), "Source": r["source"],
                   "Name in record": r["species_label"], "Country": m.get("country", ""),
                   "Metadata score (0-7)": m.get("metadata_score", ""),
                   "Missing items": m.get("metadata_missing", "").replace(";", "; "),
                   "Automated BOLD identification": m.get("identification_automated", ""),
                   "Coordinate note": m.get("coordinate_source", ""),
                   "Length (bp)": r["seq_length"], "COX1 region": r["cox1_half"],
                   "Reference start": r["ref_start"], "Reference end": r["ref_end"],
                   "Identity to reference": r["identity"], "Reading frame intact": r["reading_frame_intact"],
                   "Translation check": r.get("translation_check", ""),
                   "QC flags": r["qc_flags"].replace(";", "; "), "QC result": r["qc_status"]})
    recheck = read("results/genus/qc_recheck.csv")
    blocks2 = [("", list(s2[0].keys()), s2)]
    if recheck:
        blocks2.append(("Failed sequences and sequences with odd-length alignment gaps, rechecked against a panel of "
                        "COX1 references (one per species with a mitochondrial genome) and in six reading frames",
                        list(recheck[0].keys()), recheck))
    add_sheet(wb, "S2 Records and QC", "Table S2. Metadata and quality-control result for every record, with the "
              "translation-guided check of frame flags. Records are identified by accession and version, or by BOLD "
              "process ID for BOLD-only records.", blocks2)

    # S3: nearest-neighbour columns under accepted names (two synonymies applied)
    acc_checks = {r["record_id"]: r for r in read("results/species_accepted/record_checks.csv")}
    s3 = []
    for rid, c in checks.items():
        v, n, a = verdicts.get(rid, {}), nn.get(rid, {}), acc_checks.get(rid, c)
        s3.append({"Record": rid, "Name in record (submitted)": c["species_label"],
                   "Accepted name": v.get("accepted_name", ""),
                   "Haplotype": c["haplotype"], "Cluster (K2P 0.03)": c["cluster"],
                   "Nearest record with the same accepted name (K2P)": a["nearest_same_species"],
                   "Nearest record with another accepted name (K2P)": a["nearest_other_species"],
                   "That other accepted name": a["nearest_other_species_name"],
                   "Checks (accepted names)": a["check"].replace(";", "; "),
                   "Anchored species of cluster": v.get("anchored_species", ""), "Anchor": v.get("is_anchor", ""),
                   "Verdict": v.get("verdict", ""), "Submission batch": n.get("batch", ""),
                   "Nearest record's accepted name": n.get("nn_name_any", ""),
                   "Nearest record's accepted name, other batches": n.get("nn_name_other_batch", "")})
    add_sheet(wb, "S3 Clusters", "Table S3. Haplotype, cluster, nearest neighbours and anchor verdict for every "
              "record that passed quality control. All name comparisons use accepted names (after the two "
              "synonymies); the submitted name is given for reference.", [("", list(s3[0].keys()), s3)])

    # S4
    thr = read("results/robustness/thresholds.csv")
    grid = read("results/robustness/threshold_grid.csv")
    var = read("results/robustness/variants.csv")
    add_sheet(wb, "S4 Sensitivity", "Table S4. Sensitivity of clusters, anchor verdicts and identification "
              "accuracy to the K2P thresholds, the evidence level of anchors, the taxonomic treatment and fold "
              "isolation.",
              [("Cluster threshold and largest accepted match distance equal", list(thr[0].keys()), thr),
               ("Cluster threshold and largest accepted match distance varied separately", list(grid[0].keys()),
                grid),
               ("Anchor evidence, taxonomic treatment and fold isolation, K2P 0.03", list(var[0].keys()), var)])

    # S5
    qs = read("results/robustness/queries_main.csv")
    sp = read("results/robustness/identification_by_species.csv")
    tr = read("results/robustness/transitions.csv")
    ci = read("results/robustness/identification_intervals.csv")
    cmp_ = read("results/robustness/set_comparisons.csv")
    summ = read("results/robustness/identification_summaries.csv")
    add_sheet(wb, "S5 Identification", "Table S5. Identification of anchor queries by closest COI match "
              "(K2P 0.03, leaving out each anchor study in turn) against four reference sets: all records that "
              "passed quality control, named records, named records cleaned by anchors, and named records cleaned "
              "by the conservative rule (anchors from at least two studies).",
              [("Comparisons between reference sets: queries gained and lost, studies up and down, sign-flip test "
                "over studies and query-level McNemar test", list(cmp_[0].keys()), cmp_),
               ("Accuracy pooled, averaged over studies and over species, reference-set size, and accuracy by "
                "coverage of the query's species", list(summ[0].keys()), summ),
               ("Transitions between reference sets", list(tr[0].keys()), tr),
               ("Percentages with 95% intervals from a bootstrap over anchor studies", list(ci[0].keys()), ci),
               ("By species", list(sp[0].keys()), sp),
               ("Every query", list(qs[0].keys()), qs)])

    # S6
    dis = read("results/robustness/disagreements.csv")
    pairs = read("results/its2/coi_its2_pairs.csv")
    plab = {(r["coi_record"], r["its2_record"]): r for r in read("results/robustness/paired_label_agreement.csv")}
    for r in pairs:
        x = plab.get((r["coi_record"], r["its2_record"]), {})
        r["coi_nn_other_batch"] = x.get("coi_nn_other_batch", "")
        r["its2_nn_other_batch"] = x.get("its2_nn_other_batch", "")
    by_batch = read("results/robustness/paired_by_batch.csv") if os.path.exists(
        "results/robustness/paired_by_batch.csv") else []
    blocks = [("Records disagreeing with the anchored species of their cluster", list(dis[0].keys()), dis)]
    sp_pairs = read("results/pairs/species_pairs.csv")
    if sp_pairs:
        blocks.append(("Pairs of names sharing a cluster or within K2P 0.03: identical sequences, close but distinct, "
                       "or joined only by chaining; for all records and for anchors only", list(sp_pairs[0].keys()),
                       sp_pairs))
    if pairs:
        blocks.append(("COI and ITS2 records linked by first author and isolate, strain or voucher name (putative "
                       "material links; only one-to-one links without conflicting names or countries are used in the "
                       "comparison)", list(pairs[0].keys()), pairs))
    if by_batch:
        blocks.append(("One-to-one specimens by COI submission batch", list(by_batch[0].keys()), by_batch))
    add_sheet(wb, "S6 Pairs, disagreements, ITS2", "Table S6. Records that disagree with the anchored species of "
              "their cluster, the kind of COI overlap behind each pair of names, and the COI-ITS2 comparison for "
              "linked records.", blocks)

    # S7: every tip of the four trees (Figures 5 and S1-S4) with the records behind it,
    # in the order the tips appear in the figures (midpoint-rooted, ladderized).
    from collections import Counter
    from Bio import Phylo
    trees = [("Figure S1 (5′ half, all records)", "results/phylogeny/5prime"),
             ("Figure S2 (3′ half, all records)", "results/phylogeny/3prime"),
             ("Figure S3 (5′ half, same 98 records)", "results/phylogeny/paired/5prime"),
             ("Figure S4 (3′ half, same 98 records)", "results/phylogeny/paired/3prime")]
    s7 = []
    for label, d in trees:
        tf = os.path.join(d, "iqtree.treefile")
        if not os.path.exists(tf):
            continue
        tree = Phylo.read(tf, "newick")
        tree.root_at_midpoint()
        tree.ladderize()
        tips = {r["haplotype"]: r for r in read(os.path.join(d, "tips.csv"))}
        for pos, t in enumerate(tree.get_terminals(), start=1):
            r = tips.get(t.name, {})
            ids = (r.get("record_ids") or "").split()
            submitted = Counter(merged.get(i, {}).get("species_label", "") or "(unnamed)" for i in ids)
            s7.append({"Tree": label, "Position (top = 1)": pos, "Haplotype": t.name,
                       "Records": r.get("records", ""),
                       "Accepted names (records)": r.get("species_names", ""),
                       "Names in records as submitted (records)": "; ".join(
                           f"{k} ({v})" for k, v in submitted.most_common()),
                       "Anchor specimens": r.get("anchor_names", ""),
                       "Record identifiers (accession.version or BOLD process ID)": " ".join(ids)})
    if s7:
        add_sheet(wb, "S7 Tree tips", "Table S7. Tips of the trees in Figures 5 and S1-S4, in the order they appear "
                  "from top to bottom, with every record behind each tip.", [("", list(s7[0].keys()), s7)])

    # S8, S9: the second genus, from data/oligonychus and results/oligonychus
    G, RG = "data/oligonychus", "results/oligonychus"
    if os.path.exists(os.path.join(RG, "anchors", "anchor_status.csv")):
        o_merged = {r["record_id"]: r for r in read(f"{RG}/genus/merged_records.csv")}
        o_status = {r["accession"]: r for r in read(f"{RG}/anchors/anchor_status.csv")}
        h8, o_st, h8b, o_acc = anchor_tables(read(f"{G}/anchors/anchor_studies.csv"),
                                             read(f"{G}/anchors/anchor_records.csv"), o_status, o_merged)
        add_sheet(wb, "S8 Oligonychus anchors", "Table S8. Oligonychus: anchor sources, evidence levels and COI "
                  "accessions. Details not checked against the full text of the publication are marked unverified.",
                  [("Sources", h8, o_st), ("Anchor accessions", h8b, o_acc)])
        o_qc = read(f"{RG}/genus/qc_records.csv")
        o_checks = {r["record_id"]: r for r in read(f"{RG}/species/record_checks.csv")}
        o_ver = {r["record_id"]: r for r in read(f"{RG}/anchors/record_anchor_check.csv")}
        o_nn = {r["record_id"]: r for r in read(f"{RG}/robustness/coi_nearest_neighbour.csv")}
        recs = []
        for r in o_qc:
            m, c, v, n = (o_merged.get(r["record_id"], {}), o_checks.get(r["record_id"], {}),
                          o_ver.get(r["record_id"], {}), o_nn.get(r["record_id"], {}))
            recs.append({"Record": r["record_id"], "GenBank accession": m.get("genbank_accession", ""),
                         "BOLD process ID": m.get("bold_processid", ""), "Source": r["source"],
                         "Name in record": r["species_label"], "Country": m.get("country", ""),
                         "Metadata score (0-7)": m.get("metadata_score", ""),
                         "Missing items": m.get("metadata_missing", "").replace(";", "; "),
                         "Automated BOLD identification": m.get("identification_automated", ""),
                         "Coordinate note": m.get("coordinate_source", ""),
                         "Length (bp)": r["seq_length"], "COX1 region": r["cox1_half"],
                         "Identity to reference": r["identity"], "QC flags": r["qc_flags"].replace(";", "; "),
                         "QC result": r["qc_status"], "Haplotype": c.get("haplotype", ""),
                         "Cluster (K2P 0.03)": c.get("cluster", ""),
                         "Anchored species of cluster": v.get("anchored_species", ""),
                         "Anchor": v.get("is_anchor", ""), "Anchor used for decisions": v.get("anchor_used", ""),
                         "Verdict": v.get("verdict", ""),
                         "Submission batch": n.get("batch", "")})
        blocks = [("Every record: metadata, quality control, cluster and anchor verdict", list(recs[0].keys()), recs)]
        for heading, f in [("Identification test, every query (K2P 0.03, leaving out each anchor study in turn)",
                            "queries_main.csv"),
                           ("Comparisons between reference sets", "set_comparisons.csv"),
                           ("Accuracy pooled, averaged over studies and species, and by coverage",
                            "identification_summaries.csv"),
                           ("Identification test with the two Mexican studies held out together", "queries_grouped.csv"),
                           ("Identification test by species", "identification_by_species.csv"),
                           ("Records disagreeing with the anchored species of their cluster", "disagreements.csv"),
                           ("Cluster threshold and largest accepted match distance equal", "thresholds.csv"),
                           ("Cluster threshold and largest accepted match distance varied separately",
                            "threshold_grid.csv"),
                           ("Anchor evidence, taxonomic treatment, fold isolation and grouping, K2P 0.03", "variants.csv"),
                           ("Pairs of names sharing a cluster or within K2P 0.03", "species_pairs.csv")]:
            rows = read(f"{RG}/robustness/{f}") if not f.startswith("species_pairs") else read(f"{RG}/pairs/{f}")
            if rows:
                blocks.append((heading, list(rows[0].keys()), rows))
        rc = read(f"{RG}/genus/qc_recheck.csv")
        if rc:
            blocks.append(("Failed sequences and sequences with odd-length alignment gaps, rechecked against a panel "
                           "of references and in six reading frames", list(rc[0].keys()), rc))
        add_sheet(wb, "S9 Oligonychus results", "Table S9. Oligonychus: metadata, quality control, clusters and "
                  "anchor verdicts for every record, the identification test and its sensitivity.", blocks)

    # S10: change by submission period and country, both genera
    blocks10 = []
    for genus, d in (("Tetranychus", "results/trends"), ("Oligonychus", "results/oligonychus/trends")):
        for heading, f in (("by submission period", "by_period.csv"), ("by country", "by_country.csv")):
            rows = read(os.path.join(d, f))
            if rows:
                blocks10.append((f"{genus}, {heading}", list(rows[0].keys()), rows))
    if blocks10:
        add_sheet(wb, "S10 Change over time", "Table S10. Data quality by submission period and by country "
                  "(percentage of records), with the number of submission batches behind each country.", blocks10)

    path = os.path.join(args.out, "Supplementary_Tables_S1-S10.xlsx")
    wb.save(path)
    print(f"Written {path}")


if __name__ == "__main__":
    main()
